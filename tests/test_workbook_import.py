"""Workbook importer, using a synthetic workbook laid out like a typical hand-built budget sheet."""
from datetime import date

import openpyxl
import pytest
from openpyxl.worksheet.table import Table

from finance.db.models import Account, Bill, ImportRecord, IncomeSource
from finance.importers.budget_workbook import WorkbookImportError, import_workbook
from finance.services import accounts as acct_svc
from finance.services import dashboard

AS_OF = date(2026, 10, 24)


def build_workbook(path, *, rent=1200, car_balance=9000.55, checking=250):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A5"], ws["B5"], ws["C5"] = "Income", "Amount", "Type"
    ws["A6"], ws["B6"], ws["C6"] = "Acme Corp (Income)", 3000, "Full Time Job"
    ws["A7"], ws["B7"], ws["C7"] = "Bills (Owe)", "=SUM('Monthly Expenses'!B3:B8)", "Monthly Dues"
    ws["A8"], ws["B8"] = "Total ", "=SUM(B6-B7)"
    ws.add_table(Table(displayName="Monthly_", ref="A5:C8"))

    ws["A15"], ws["B15"], ws["C15"] = "Person", "Amount", "Type"
    ws["A16"], ws["B16"], ws["C16"] = "Pat", 0, "Help"
    # row 17 intentionally blank
    ws["A18"], ws["B18"], ws["C18"] = "Sam", 75, "Concert"
    ws["A19"], ws["B19"] = "Total", "=SUM(B16:B18)"

    ws["A27"], ws["B27"], ws["C27"] = "Type ", "Allowed Amount", "Owe"
    ws["A28"], ws["B28"], ws["C28"] = "Visa", 5000, 410.25
    ws["A29"], ws["B29"] = "Total", "=SUM(B28:B28)"

    ws["N35"], ws["O35"] = "Type ", "Amount"
    for i, (n, v) in enumerate([("Checkings", checking), ("Savings ", 4000), ("Gift card", 25), ("Retirement", 1500.5)]):
        ws.cell(36 + i, 14, n)
        ws.cell(36 + i, 15, v)
    ws.cell(40, 14, "Total ")
    ws.cell(40, 15, "=SUM(O36:O39)")

    ex = wb.create_sheet("Monthly Expenses")
    ex.append([])
    ex.append(["Expenses", "Amount", "Type", "Aiming to Pay Monthly", "Monthly Date"])
    for row in [
        ["Car ", 350, car_balance, None, "15th"],
        ["Rent", rent, None, None, "1st"],
        ["Streaming", 15, None, None, None],
        ["Roth IRA", 50, None, None, None],
        ["Mystery Thing", 20, None, None, None],
        ["Groceries", 300, None, None, None],
    ]:
        ex.append(row)
    ex.append(["Total", "=SUM(B3:B8)", "=SUM(C3:C8)"])
    wb.save(path)
    return path


def test_full_import(session, tmp_path):
    p = build_workbook(tmp_path / "budget.xlsx")
    res = import_workbook(session, p, as_of=AS_OF)
    session.flush()

    accts = {(a.name, a.account_type): a for a in session.query(Account)}
    assert accts[("Checkings", "checking")].current_balance_cents == 250_00
    assert accts[("Savings", "savings")].current_balance_cents == 4000_00
    assert accts[("Retirement", "retirement")].current_balance_cents == 1500_50
    gift = accts[("Gift card", "other_asset")]
    assert gift.needs_review  # unknown label -> flagged, not guessed
    visa = accts[("Visa", "credit_card")]
    assert (visa.current_balance_cents, visa.credit_limit_cents) == (410_25, 5000_00)
    car = accts[("Car loan", "auto_loan")]
    assert (car.current_balance_cents, car.minimum_payment_cents, car.payment_due_day) == (9000_55, 350_00, 15)
    assert accts[("Sam", "receivable")].current_balance_cents == 75_00
    assert ("Pat", "receivable") not in accts  # $0 rows skipped and reported
    assert any("Pat" in s for s in res.skipped)

    bills = {b.name: b for b in session.query(Bill)}
    assert set(bills) == {"Car", "Rent", "Streaming", "Roth IRA", "Mystery Thing", "Groceries"}
    assert bills["Car"].kind == "debt_payment" and bills["Car"].pays_account_id == car.id
    assert bills["Rent"].due_day == 1 and bills["Rent"].kind == "bill"
    assert bills["Roth IRA"].kind == "savings"
    assert bills["Groceries"].kind == "planned"
    assert bills["Mystery Thing"].needs_review
    assert bills["Streaming"].needs_review and "due date" in bills["Streaming"].review_note
    assert bills["Rent"].tracking_start == AS_OF

    inc = session.query(IncomeSource).one()
    assert (inc.name, inc.amount_cents, inc.kind) == ("Acme Corp", 3000_00, "salary")
    assert any("Bills" in s for s in res.skipped)  # formula row in income table not imported

    d = dashboard.build(session, today=AS_OF)
    assert d.income_cents == 3000_00
    assert d.expenses_cents == (350 + 1200 + 15 + 20 + 300) * 100
    assert d.savings_contrib_cents == 50_00
    assert d.left_after_plan_cents == (3000 - 1935) * 100
    assert d.assets_cents == (250 + 4000 + 25 + 1500.5 + 75) * 100
    assert d.liabilities_cents == 410_25 + 9000_55
    assert session.query(ImportRecord).count() == 1


def test_same_file_twice_is_refused(session, tmp_path):
    p = build_workbook(tmp_path / "budget.xlsx")
    import_workbook(session, p, as_of=AS_OF)
    again = import_workbook(session, p, as_of=AS_OF)
    assert again.already_imported and again.total_changes == 0
    assert session.query(Account).count() == 7


def test_reimport_updates_without_deleting(session, tmp_path):
    import_workbook(session, build_workbook(tmp_path / "v1.xlsx"), as_of=date(2026, 10, 1))
    session.flush()
    res = import_workbook(session, build_workbook(tmp_path / "v2.xlsx", rent=1250, car_balance=8650.55, checking=900),
                          as_of=AS_OF)
    session.flush()
    rent = session.query(Bill).filter_by(name="Rent").one()
    assert rent.amount_cents == 1250_00
    car = acct_svc.find_by_name(session, "Car loan", "auto_loan")
    assert car.current_balance_cents == 8650_55
    assert [h.balance_cents for h in acct_svc.history(session, car.id)] == [9000_55, 8650_55]
    assert any("Rent" in u for u in res.updated)
    assert session.query(Bill).count() == 6 and session.query(Account).count() == 7


def test_reimport_keeps_user_category_edits(session, tmp_path):
    import_workbook(session, build_workbook(tmp_path / "v1.xlsx"), as_of=AS_OF)
    session.flush()
    m = session.query(Bill).filter_by(name="Mystery Thing").one()
    m.kind = "planned"
    import_workbook(session, build_workbook(tmp_path / "v2.xlsx", rent=1300), as_of=AS_OF)
    assert session.query(Bill).filter_by(name="Mystery Thing").one().kind == "planned"


def test_unrecognised_workbook(session, tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "hello"
    wb.save(tmp_path / "x.xlsx")
    with pytest.raises(WorkbookImportError, match="No recognisable"):
        import_workbook(session, tmp_path / "x.xlsx")


def test_not_a_workbook(session, tmp_path):
    (tmp_path / "bad.xlsx").write_bytes(b"not a zip")
    with pytest.raises(WorkbookImportError):
        import_workbook(session, tmp_path / "bad.xlsx")
